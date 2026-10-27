import json, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
D=json.load(open('products_final.json')); P=D['products']
wb=openpyxl.Workbook()

rm=wb.active; rm.title='Read me'
rows=[('CARPENTA PRODUCT CATALOGUE — MASTER DATA',''),('',''),
 ('This workbook drives the catalogue. Edit here, then ask Claude to rebuild.',''),('',''),
 ('YELLOW columns are yours to change','Show, Product name, Price (GHS), Note'),
 ('Show','Y shows the product in the catalogue. Anything else hides it.'),
 ('Price (GHS)','Numbers only, no currency symbol or commas.'),
 ('Photo link','The image used. Blank means designed artwork is used instead.'),
 ('Image type','How the picture is produced — see below.'),('',''),
 ('Photo','A real photograph already on carpenta.store.'),
 ('Photo (own)','A photograph you supplied. The file sits in the photos folder, named by SKU.'),
 ('Photo (manufacturer)','Sourced from the brand’s own website.'),
 ('Roll — your decor photo','Drawn as a real edge-band roll using your board photo as the decor.'),
 ('Roll — colour from your roll photo','Same roll, coloured from the roll photographs you supplied.'),
 ('Colour guide','A flat colour swatch, used for paints.'),
 ('Designed tile','No picture yet. Still on the photo shoot list.'),('',''),
 ('Rebuild','python3 src/recover.py  then  python3 src/build_catalogue2.py')]
for r in rows: rm.append(r)
rm['A1'].font=Font(size=14,bold=True); rm.column_dimensions['A'].width=34; rm.column_dimensions['B'].width=74
for i in range(6,18): rm.cell(i,1).font=Font(bold=True)

ws=wb.create_sheet('Products')
hdr=['Show','Category','Group','Product name','SKU','Price (GHS)','Photo link','Image type','Note']
ws.append(hdr)
YEL=PatternFill('solid',fgColor='FFF6D5'); HDR=PatternFill('solid',fgColor='0B1016')
thin=Side(style='thin',color='D4D8DE')
for c in range(1,len(hdr)+1):
    x=ws.cell(1,c); x.font=Font(bold=True,color='FFFFFF',size=10); x.fill=HDR
    x.alignment=Alignment(vertical='center'); x.border=Border(bottom=thin)
ws.freeze_panes='A2'
for p in P:
    ws.append([p['show'],p['cat'],p['sub'],p['name'],p['sku'],p['price'],
               p.get('local') or p['img'],p['itype'],p.get('note','')])
for r in range(2,len(P)+2):
    for c in (1,4,6,9): ws.cell(r,c).fill=YEL
    ws.cell(r,6).number_format='#,##0'
    for c in range(1,10): ws.cell(r,c).border=Border(bottom=thin)
for col,w in zip('ABCDEFGHI',[6,20,20,52,16,12,58,30,30]): ws.column_dimensions[col].width=w
ws.auto_filter.ref=f'A1:I{len(P)+1}'

sm=wb.create_sheet('Summary')
from collections import Counter
sm.append(['Image type','Count']); sm['A1'].font=Font(bold=True); sm['B1'].font=Font(bold=True)
for k,v in Counter(p['itype'] for p in P).most_common(): sm.append([k,v])
sm.append(['','']); sm.append(['Total',len(P)]); sm.cell(sm.max_row,1).font=Font(bold=True)
sm.column_dimensions['A'].width=36; sm.column_dimensions['B'].width=10
wb.save('../output/Carpenta_Products.xlsx')
print('xlsx', len(P),'rows')
